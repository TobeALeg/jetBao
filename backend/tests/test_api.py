from __future__ import annotations

import importlib
import json
import sys
import zipfile
from io import BytesIO
from pathlib import Path
from urllib.parse import parse_qs, urlparse

from fastapi.testclient import TestClient
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))


def make_client(
    tmp_path: Path,
    monkeypatch,
    seed_demo_users: str = "true",
    bootstrap_admin: dict[str, str] | None = None,
    auth_mode: str = "hybrid",
) -> TestClient:
    monkeypatch.setenv("DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("UPLOAD_DIR", str(tmp_path / "data" / "uploads"))
    monkeypatch.setenv("SECRET_KEY", "test-secret")
    monkeypatch.setenv("SEED_DEMO_USERS", seed_demo_users)
    monkeypatch.setenv("TENCENT_SECRET_ID", "")
    monkeypatch.setenv("TENCENT_SECRET_KEY", "")
    monkeypatch.setenv("AUTH_MODE", auth_mode)
    monkeypatch.setenv("SSO_AUTHORIZE_URL", "https://mentti.work/sso/authorize")
    monkeypatch.setenv("SSO_TOKEN_URL", "https://mentti.work/api/sso/token")
    monkeypatch.setenv("SSO_CLIENT_ID", "jetbao")
    monkeypatch.setenv("SSO_CLIENT_SECRET", "test-client-secret")
    monkeypatch.setenv("SSO_REDIRECT_URI", "https://jetbao.mentti.work/api/auth/sso/callback")
    monkeypatch.setenv("SSO_COOKIE_SECURE", "false")
    if bootstrap_admin:
        for key, value in bootstrap_admin.items():
            monkeypatch.setenv(key, value)
    main = importlib.import_module("app.main")
    app = main.create_app()
    app.state.db.init(app.state.settings.seed_demo_users, app.state.settings.bootstrap_admin)
    return TestClient(app)


def auth_headers(client: TestClient, username: str, password: str) -> dict[str, str]:
    response = client.post("/api/auth/login", json={"username": username, "password": password})
    assert response.status_code == 200
    token = response.json()["token"]
    return {"Authorization": f"Bearer {token}"}


def provision_email(client: TestClient, username: str, email: str) -> None:
    with client.app.state.db.connect() as connection:
        connection.execute(
            "UPDATE users SET email = ? WHERE username = ?",
            (email, username),
        )


class StubSsoClient:
    def __init__(self, identity: dict[str, str]):
        self.identity = identity

    def exchange_code(self, code: str) -> dict[str, str]:
        assert code == "single-use-code"
        return self.identity


def upload_file(
    client: TestClient,
    headers: dict[str, str],
    filename: str,
    content: bytes,
    content_type: str = "application/pdf",
):
    return client.post(
        "/api/attachments",
        headers=headers,
        files={"file": (filename, content, content_type)},
    )


def upload_evidence_file(
    client: TestClient,
    headers: dict[str, str],
    filename: str,
    content: bytes,
    content_type: str = "application/pdf",
):
    return client.post(
        "/api/attachments/batch",
        headers=headers,
        files=[("files", (filename, content, content_type))],
    )


def insert_ocr_attachment(
    client: TestClient,
    user_id: int,
    invoice_items: list[dict],
    filename: str = "ocr.pdf",
    pool_status: str = "pooled",
) -> int:
    stored_path = client.app.state.settings.upload_dir / filename
    stored_path.parent.mkdir(parents=True, exist_ok=True)
    stored_path.write_bytes(f"test file for {filename}".encode())
    with client.app.state.db.connect() as connection:
        cursor = connection.execute(
            """
            INSERT INTO attachments (
                user_id, original_filename, stored_path, file_hash, file_size,
                duplicate_count, pool_status, ocr_status, ocr_result
            )
            VALUES (?, ?, ?, ?, ?, 0, ?, 'success', ?)
            """,
            (
                user_id,
                filename,
                str(stored_path),
                f"hash-{filename}",
                stored_path.stat().st_size,
                pool_status,
                json.dumps({"invoice_items": invoice_items}, ensure_ascii=False),
            ),
        )
        return int(cursor.lastrowid)


def create_invoice_backed_expense(
    client: TestClient,
    headers: dict[str, str],
    user_id: int,
    amount: float = 300,
    project_name: str = "市场活动",
    category: str = "市场活动",
    month: str = "2026-05",
    invoice_number: str = "INV-OK",
    buyer: str = "上海山途远智信息科技有限公司",
    buyer_confirmed: bool = False,
) -> dict:
    attachment_id = insert_ocr_attachment(
        client,
        user_id=user_id,
        invoice_items=[
            {
                "buyer": buyer,
                "amount": amount,
                "invoice_number": invoice_number,
                "date": "2026-05-20",
                "sub_type_description": "电子普通发票",
            }
        ],
        filename=f"{invoice_number}.pdf",
    )
    draft = client.post(
        "/api/expenses/drafts",
        headers=headers,
        json={
            "project_name": project_name,
            "actual_amount": amount,
            "expense_month": month,
            "category": category,
        },
    )
    assert draft.status_code == 200
    match = client.post(
        "/api/expense-allocations",
        headers=headers,
        json={
            "expense_id": draft.json()["id"],
            "attachment_id": attachment_id,
            "invoice_item_index": 0,
            "note": "",
            "buyer_confirmed": buyer_confirmed,
        },
    )
    assert match.status_code == 200
    return match.json()


def test_dandi_and_ouyang_are_seeded_as_admins(tmp_path, monkeypatch):
    client = make_client(tmp_path, monkeypatch)

    for username, password, employee_name in (
        ("Dandi", "dandi123", "艾丹迪"),
        ("Ouyang", "ouyang123", "欧阳"),
    ):
        headers = auth_headers(client, username, password)
        me = client.get("/api/me", headers=headers)
        assert me.status_code == 200
        assert me.json()["role"] == "admin"
        assert me.json()["employee_name"] == employee_name


def test_demo_users_are_not_seeded_by_default(tmp_path, monkeypatch):
    client = make_client(tmp_path, monkeypatch, seed_demo_users="false")

    response = client.post("/api/auth/login", json={"username": "admin", "password": "admin123"})

    assert response.status_code == 401


def test_bootstrap_admin_creates_first_admin_without_demo_users(tmp_path, monkeypatch):
    client = make_client(
        tmp_path,
        monkeypatch,
        seed_demo_users="false",
        bootstrap_admin={
            "BOOTSTRAP_ADMIN_USERNAME": "owner",
            "BOOTSTRAP_ADMIN_PASSWORD": "owner-pass",
            "BOOTSTRAP_ADMIN_EMPLOYEE_NAME": "Owner",
            "BOOTSTRAP_ADMIN_COMPANY_ENTITY": "上海山途远智信息科技有限公司",
        },
    )

    headers = auth_headers(client, "owner", "owner-pass")
    me = client.get("/api/me", headers=headers)

    assert me.status_code == 200
    assert me.json()["role"] == "admin"
    assert me.json()["employee_name"] == "Owner"


def test_sso_login_uses_preprovisioned_email_and_creates_cookie_session(tmp_path, monkeypatch):
    client = make_client(tmp_path, monkeypatch, auth_mode="sso")
    provision_email(client, "Dandi", "dandi@mentitrek.com")
    client.app.state.sso_client = StubSsoClient(
        {"sub": "mentti-user-2", "email": "dandi@mentitrek.com"}
    )

    start = client.get("/api/auth/sso/start", follow_redirects=False)

    assert start.status_code == 307
    location = urlparse(start.headers["location"])
    query = parse_qs(location.query)
    assert f"{location.scheme}://{location.netloc}{location.path}" == "https://mentti.work/sso/authorize"
    assert query == {
        "client_id": ["jetbao"],
        "redirect_uri": ["https://jetbao.mentti.work/api/auth/sso/callback"],
        "state": [query["state"][0]],
    }
    assert start.cookies["jetbao_sso_state"] == query["state"][0]

    callback = client.get(
        "/api/auth/sso/callback",
        params={"code": "single-use-code", "state": query["state"][0]},
        follow_redirects=False,
    )

    assert callback.status_code == 303
    assert callback.headers["location"] == "/"
    assert "HttpOnly" in callback.headers["set-cookie"]
    me = client.get("/api/me")
    assert me.status_code == 200
    assert me.json()["email"] == "dandi@mentitrek.com"
    assert me.json()["role"] == "admin"
    with client.app.state.db.connect() as connection:
        user = connection.execute("SELECT identity_id FROM users WHERE username = 'Dandi'").fetchone()
    assert user["identity_id"] == "mentti-user-2"


def test_sso_login_rejects_enterprise_email_without_jetbao_access(tmp_path, monkeypatch):
    client = make_client(tmp_path, monkeypatch, auth_mode="sso")
    client.app.state.sso_client = StubSsoClient(
        {"sub": "mentti-new-user", "email": "new.employee@mentitrek.com"}
    )
    start = client.get("/api/auth/sso/start", follow_redirects=False)
    state = parse_qs(urlparse(start.headers["location"]).query)["state"][0]

    callback = client.get(
        "/api/auth/sso/callback",
        params={"code": "single-use-code", "state": state},
        follow_redirects=False,
    )

    assert callback.status_code == 303
    assert callback.headers["location"].startswith("/?sso_error=")
    assert client.get("/api/me").status_code == 401


def test_sso_callback_rejects_invalid_state_before_code_exchange(tmp_path, monkeypatch):
    client = make_client(tmp_path, monkeypatch, auth_mode="sso")
    client.app.state.sso_client = StubSsoClient(
        {"sub": "mentti-user-2", "email": "dandi@mentitrek.com"}
    )
    client.get("/api/auth/sso/start", follow_redirects=False)

    callback = client.get(
        "/api/auth/sso/callback",
        params={"code": "single-use-code", "state": "tampered"},
        follow_redirects=False,
    )

    assert callback.status_code == 400


def test_dandi_username_is_not_auto_promoted_on_restart(tmp_path, monkeypatch):
    client = make_client(
        tmp_path,
        monkeypatch,
        seed_demo_users="false",
        bootstrap_admin={
            "BOOTSTRAP_ADMIN_USERNAME": "owner",
            "BOOTSTRAP_ADMIN_PASSWORD": "owner-pass",
            "BOOTSTRAP_ADMIN_EMPLOYEE_NAME": "Owner",
            "BOOTSTRAP_ADMIN_COMPANY_ENTITY": "上海山途远智信息科技有限公司",
        },
    )
    owner = auth_headers(client, "owner", "owner-pass")
    created = client.post(
        "/api/admin/users",
        headers=owner,
        json={
            "username": "Dandi",
            "password": "dandi123",
            "role": "employee",
            "employee_name": "艾丹迪",
            "company_entity": "上海山途远智信息科技有限公司",
        },
    )
    assert created.status_code == 200

    client.app.state.db.init(seed_demo_users=False, bootstrap_admin=None)
    users = client.get("/api/admin/users", headers=owner)
    dandi = next(user for user in users.json() if user["username"] == "Dandi")

    assert dandi["role"] == "employee"


def test_legacy_alice_and_bob_demo_accounts_are_not_seeded(tmp_path, monkeypatch):
    client = make_client(tmp_path, monkeypatch)

    for username, password in (("alice", "alice123"), ("bob", "bob123")):
        response = client.post("/api/auth/login", json={"username": username, "password": password})
        assert response.status_code == 401


def test_user_can_change_own_password(tmp_path, monkeypatch):
    client = make_client(tmp_path, monkeypatch)
    headers = auth_headers(client, "Dandi", "dandi123")

    wrong_current = client.patch(
        "/api/me/password",
        headers=headers,
        json={"current_password": "wrong-password", "new_password": "dandi-new123"},
    )
    assert wrong_current.status_code == 400

    changed = client.patch(
        "/api/me/password",
        headers=headers,
        json={"current_password": "dandi123", "new_password": "dandi-new123"},
    )
    assert changed.status_code == 200
    assert changed.json()["username"] == "Dandi"

    old_login = client.post("/api/auth/login", json={"username": "Dandi", "password": "dandi123"})
    assert old_login.status_code == 401
    new_login = client.post("/api/auth/login", json={"username": "Dandi", "password": "dandi-new123"})
    assert new_login.status_code == 200


def test_employee_can_create_draft_and_only_see_own_records(tmp_path, monkeypatch):
    client = make_client(tmp_path, monkeypatch)
    dandi = auth_headers(client, "Dandi", "dandi123")
    ouyang = auth_headers(client, "Ouyang", "ouyang123")

    create = client.post(
        "/api/expenses/drafts",
        headers=dandi,
        json={
            "project_name": "机场快线",
            "category": "差旅交通",
            "expense_month": "2026-05",
            "actual_amount": 128.5,
        },
    )
    assert create.status_code == 200
    assert create.json()["company_entity"] == "上海山途远智信息科技有限公司"
    assert create.json()["status"] == "pending"

    dandi_records = client.get("/api/expenses", headers=dandi)
    ouyang_records = client.get("/api/expenses", headers=ouyang)
    assert len(dandi_records.json()) == 1
    assert ouyang_records.json() == []


def test_expense_endpoint_creates_pending_record(tmp_path, monkeypatch):
    client = make_client(tmp_path, monkeypatch)
    dandi = auth_headers(client, "Dandi", "dandi123")

    response = client.post(
        "/api/expenses",
        headers=dandi,
        json={
            "project_name": "AI 工具订阅",
            "category": "AI 项目",
            "expense_month": "2026-05",
            "actual_amount": 100,
            "is_substitute": False,
            "substitute_reason": "",
        },
    )
    assert response.status_code == 200
    assert response.json()["status"] == "pending"


def test_duplicate_upload_is_marked_but_not_blocked(tmp_path, monkeypatch):
    client = make_client(tmp_path, monkeypatch)
    dandi = auth_headers(client, "Dandi", "dandi123")

    first = upload_evidence_file(client, dandi, "a.pdf", b"same-content")
    second = upload_evidence_file(client, dandi, "b.pdf", b"same-content")
    assert first.status_code == 200
    assert second.status_code == 200
    assert first.json()[0]["is_duplicate"] is False
    assert second.json()[0]["is_duplicate"] is True


def test_admin_can_filter_ledger_and_preview_export(tmp_path, monkeypatch):
    client = make_client(tmp_path, monkeypatch)
    dandi = auth_headers(client, "Dandi", "dandi123")
    admin = auth_headers(client, "admin", "admin123")

    create = create_invoice_backed_expense(client, dandi, user_id=2, amount=300)
    submitted = client.post(f"/api/expenses/{create['id']}/submit", headers=dandi)
    assert submitted.status_code == 200
    assert submitted.json()["status"] == "matched"

    draft = client.post(
        "/api/expenses/drafts",
        headers=dandi,
        json={
            "project_name": "客户拜访打车",
            "actual_amount": 80,
            "expense_month": "2026-05",
            "category": "差旅交通",
        },
    )
    assert draft.status_code == 200

    ledger = client.get("/api/admin/ledger?month=2026-05&employee=艾丹迪&status=matched", headers=admin)
    assert ledger.status_code == 200
    assert len(ledger.json()) == 1
    assert ledger.json()[0]["employee_name"] == "艾丹迪"
    assert ledger.json()[0]["status"] == "matched"

    draft_ledger = client.get("/api/admin/ledger?month=2026-05&status=pending", headers=admin)
    assert draft_ledger.status_code == 200
    assert len(draft_ledger.json()) == 1
    assert draft_ledger.json()[0]["project_name"] == "客户拜访打车"

    preview = client.get("/api/admin/export/preview?month=2026-05", headers=admin)
    assert preview.status_code == 200
    assert preview.json() == {
        "employee_count": 1,
        "record_count": 1,
        "total_amount": 300.0,
        "pending_count": 1,
    }

    export = client.get("/api/admin/export.xlsx?month=2026-05", headers=admin)
    assert export.status_code == 200
    assert export.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


def test_admin_can_export_detail_package_with_workbook_and_files(tmp_path, monkeypatch):
    client = make_client(tmp_path, monkeypatch)
    dandi = auth_headers(client, "Dandi", "dandi123")
    ouyang = auth_headers(client, "Ouyang", "ouyang123")
    admin = auth_headers(client, "admin", "admin123")

    payment = upload_evidence_file(client, dandi, "payment.png", b"payment-image", "image/png")
    assert payment.status_code == 200
    dandi_invoice_id = insert_ocr_attachment(
        client,
        user_id=2,
        invoice_items=[
            {
                "buyer": "上海山途远智信息科技有限公司",
                "seller_name": "上海出租车公司",
                "item_name": "出租车费",
                "amount": 420,
                "invoice_number": "INV-A",
                "date": "2026-05-21",
                "sub_type_description": "电子普通发票",
            }
        ],
        filename="invoice-a.pdf",
    )
    draft = client.post(
        "/api/expenses/drafts",
        headers=dandi,
        json={
            "project_name": "客户拜访差旅",
            "actual_amount": 420,
            "expense_month": "2026-05",
            "category": "差旅交通",
        },
    )
    assert draft.status_code == 200
    link = client.post(
        f"/api/expenses/{draft.json()['id']}/attachments",
        headers=dandi,
        json={"attachment_ids": [payment.json()[0]["id"]]},
    )
    assert link.status_code == 200
    match = client.post(
        "/api/expense-allocations",
        headers=dandi,
        json={
            "expense_id": draft.json()["id"],
            "attachment_id": dandi_invoice_id,
            "invoice_item_index": 0,
            "note": "",
        },
    )
    assert match.status_code == 200, match.text

    ouyang_invoice_id = insert_ocr_attachment(
        client,
        user_id=3,
        invoice_items=[
            {
                "buyer": "上海山途远智信息科技有限公司",
                "seller_name": "杭州办公用品有限公司",
                "item_name": "办公耗材",
                "amount": 80,
                "invoice_number": "OY-1",
                "date": "2026-05-23",
                "sub_type_description": "电子普通发票",
            }
        ],
        filename="ouyang-office.pdf",
    )
    ouyang_draft = client.post(
        "/api/expenses/drafts",
        headers=ouyang,
        json={
            "project_name": "办公耗材",
            "actual_amount": 80,
            "expense_month": "2026-05",
            "category": "办公采购",
        },
    )
    assert ouyang_draft.status_code == 200
    ouyang_match = client.post(
        "/api/expense-allocations",
        headers=ouyang,
        json={
            "expense_id": ouyang_draft.json()["id"],
            "attachment_id": ouyang_invoice_id,
            "invoice_item_index": 0,
            "note": "",
        },
    )
    assert ouyang_match.status_code == 200

    dandi_submit = client.post(f"/api/expenses/{draft.json()['id']}/submit", headers=dandi)
    assert dandi_submit.status_code == 200
    ouyang_submit = client.post(f"/api/expenses/{ouyang_draft.json()['id']}/submit", headers=ouyang)
    assert ouyang_submit.status_code == 200

    package = client.get("/api/admin/export-package.zip?month=2026-05", headers=admin)
    assert package.status_code == 200
    assert package.headers["content-type"].startswith("application/zip")

    archive = zipfile.ZipFile(BytesIO(package.content))
    names = archive.namelist()
    assert "5月报销明细.xlsx" in names
    assert any(name.startswith("山途远智5月报销/差旅交通05月报销/艾丹迪05月报销/") for name in names)
    assert "山途远智5月报销/差旅交通05月报销/艾丹迪05月报销/客户拜访差旅佐证材料420元.png" in names
    assert "山途远智5月报销/差旅交通05月报销/艾丹迪05月报销/客户拜访差旅发票420元.pdf" in names
    assert "山途远智5月报销/办公采购05月报销/欧阳05月报销/办公耗材发票80元.pdf" in names

    workbook = load_workbook(BytesIO(archive.read("5月报销明细.xlsx")))
    assert workbook.sheetnames[:4] == ["总览", "报销项汇总", "发票明细", "附件与待核对"]
    assert "_chart_data" in workbook.sheetnames
    overview = workbook["总览"]
    assert len(overview._charts) == 1
    assert overview._charts[0].title.tx.rich.p[0].r[0].t == "按人员本次报销金额"
    assert overview["A1"].value == "5月发票/报销整理总览"
    assert overview["A4"].value == "本次报销金额合计"
    assert overview["A5"].value == 500
    assert overview["C4"].value == "票面金额合计"
    assert overview["C5"].value == 500
    assert overview["E4"].value == "涉及人员"
    assert overview["E5"].value == 2
    assert overview["G4"].value == "发票张数"
    assert overview["G5"].value == 2
    assert [overview.cell(row=7, column=column).value for column in range(1, 8)] == [
        "公司主体",
        "人员数",
        "报销项数",
        "发票张数",
        "票面金额合计",
        "本次报销金额",
        "票面-报销差异",
    ]
    detail_header_row = next(
        row
        for row in range(1, overview.max_row + 1)
        if overview.cell(row=row, column=1).value == "公司主体" and overview.cell(row=row, column=2).value == "人员"
    )
    company_rows = {
        overview.cell(row=row, column=1).value: [overview.cell(row=row, column=column).value for column in range(2, 8)]
        for row in range(8, detail_header_row)
    }
    assert company_rows["信息科技"] == [2, 2, 2, 500, 500, 0]
    detail_rows = {
        overview.cell(row=row, column=2).value: [overview.cell(row=row, column=column).value for column in range(1, 8)]
        for row in range(detail_header_row + 1, detail_header_row + 3)
    }
    assert detail_rows["艾丹迪"] == ["信息科技", "艾丹迪", 1, 1, 420, 420, 0]
    assert detail_rows["欧阳"] == ["信息科技", "欧阳", 1, 1, 80, 80, 0]
    summary = workbook["报销项汇总"]
    assert [cell.value for cell in summary[1]] == [
        "组ID",
        "公司主体",
        "购买方名称",
        "人员",
        "报销项",
        "票据状态",
        "发票张数",
        "票面金额合计",
        "本次报销金额",
        "票面-报销差异",
        "附件数",
        "备注",
    ]
    summary_rows = {
        summary.cell(row=row, column=5).value: [
            summary.cell(row=row, column=4).value,
            *[summary.cell(row=row, column=column).value for column in range(6, 10)],
        ]
        for row in range(2, summary.max_row + 1)
    }
    # 报销项汇总：按人员 + 报销类别归类，报销项列填写类别本身
    assert summary_rows["差旅交通"] == ["艾丹迪", "正常发票/行程单", 1, 420, 420]
    assert summary_rows["办公采购"] == ["欧阳", "正常发票", 1, 80, 80]
    assert [summary.cell(row=row, column=1).value for row in range(2, summary.max_row + 1)] == ["G01", "G02"]
    assert {summary.cell(row=row, column=12).value for row in range(2, summary.max_row + 1)} == {"等待通过"}
    assert all(summary.cell(row=row, column=12).font.color.rgb.endswith("FF0000") for row in range(2, summary.max_row + 1))

    invoice_sheet = workbook["发票明细"]
    assert [cell.value for cell in invoice_sheet[1]] == [
        "公司主体",
        "购买方名称",
        "人员",
        "编号",
        "报销事项",
        "发票/单据类型",
        "发票号码",
        "开票日期",
        "销售方",
        "项目名称",
        "报销价税合计",
        "是否查验",
        "发票文件路径",
        "备注",
    ]
    assert invoice_sheet.max_row == 3
    assert {invoice_sheet[f"G{row}"].value for row in range(2, 4)} == {"INV-A", "OY-1"}
    attachments_sheet = workbook["附件与待核对"]
    attachment_types = {attachments_sheet[f"D{row}"].value for row in range(2, attachments_sheet.max_row + 1)}
    attachment_names = {attachments_sheet[f"E{row}"].value for row in range(2, attachments_sheet.max_row + 1)}
    attachment_paths = {attachments_sheet[f"G{row}"].value for row in range(2, attachments_sheet.max_row + 1)}
    attachment_companies = {attachments_sheet[f"B{row}"].value for row in range(2, attachments_sheet.max_row + 1)}
    assert attachment_types == {"发票图片"}
    assert attachment_names == {"客户拜访差旅发票420元.pdf", "办公耗材发票80元.pdf"}
    assert all("发票" in str(name) and str(name).endswith("元.pdf") for name in attachment_names)
    assert all("佐证材料" not in str(path) for path in attachment_paths)
    assert all(str(path).endswith("元.pdf") for path in attachment_paths)
    # 附件页公司主体取个人资料绑定的完整公司名，不用简称
    assert attachment_companies == {"上海山途远智信息科技有限公司"}


def test_employee_can_create_draft_and_complete_it_with_invoice_item(tmp_path, monkeypatch):
    client = make_client(tmp_path, monkeypatch)
    dandi = auth_headers(client, "Dandi", "dandi123")
    attachment_id = insert_ocr_attachment(
        client,
        user_id=2,
        invoice_items=[
            {
                "buyer": "上海山途远智信息科技有限公司",
                "amount": 120,
                "invoice_number": "DRAFT-1",
                "date": "2026-05-20",
                "sub_type_description": "出租车票",
            }
        ],
        filename="draft-proof.pdf",
    )

    draft = client.post(
        "/api/expenses/drafts",
        headers=dandi,
        json={
            "project_name": "客户拜访打车",
            "actual_amount": 120,
            "expense_month": "2026-05",
            "category": "差旅交通",
        },
    )
    assert draft.status_code == 200
    body = draft.json()
    assert body["status"] == "pending"
    assert body["project_name"] == "客户拜访打车"

    match = client.post(
        "/api/expense-allocations",
        headers=dandi,
        json={
            "expense_id": body["id"],
            "attachment_id": attachment_id,
            "invoice_item_index": 0,
            "note": "拜访 A 客户",
        },
    )
    assert match.status_code == 200
    matched = match.json()
    assert matched["status"] == "pending"
    submit = client.post(f"/api/expenses/{body['id']}/submit", headers=dandi)
    assert submit.status_code == 200
    completed = submit.json()
    assert completed["status"] == "matched"
    assert completed["project_name"] == "客户拜访打车"
    assert completed["invoice_amount"] == 120
    assert completed["invoice_number"] == "DRAFT-1"
    assert len(completed["attachments"]) == 0
    assert len(completed["allocations"]) == 1
    assert completed["allocations"][0]["attachment_id"] == attachment_id


def test_one_expense_can_match_multiple_invoices_but_invoice_is_single_owner(tmp_path, monkeypatch):
    client = make_client(tmp_path, monkeypatch)
    dandi = auth_headers(client, "Dandi", "dandi123")
    first_attachment_id = insert_ocr_attachment(
        client,
        user_id=2,
        invoice_items=[
            {
                "buyer": "上海山途远智信息科技有限公司",
                "amount": 120,
                "invoice_number": "INV-1",
                "date": "2026-05-21",
                "sub_type_description": "电子普通发票",
            }
        ],
        filename="first-invoice.pdf",
    )
    second_attachment_id = insert_ocr_attachment(
        client,
        user_id=2,
        invoice_items=[
            {
                "buyer": "上海山途远智信息科技有限公司",
                "amount": 300,
                "invoice_number": "INV-2",
                "date": "2026-05-22",
                "sub_type_description": "电子普通发票",
            }
        ],
        filename="second-invoice.pdf",
    )

    expense = client.post(
        "/api/expenses/drafts",
        headers=dandi,
        json={
            "project_name": "客户拜访差旅",
            "actual_amount": 420,
            "expense_month": "2026-05",
            "category": "差旅交通",
        },
    ).json()

    match = client.post(
        "/api/expense-allocations/batch",
        headers=dandi,
        json={
            "expense_id": expense["id"],
            "invoices": [
                {"attachment_id": first_attachment_id, "invoice_item_index": 0},
                {"attachment_id": second_attachment_id, "invoice_item_index": 0},
            ],
            "note": "",
        },
    )
    assert match.status_code == 200
    matched = match.json()
    assert matched["status"] == "pending"
    assert matched["invoice_amount"] == 420
    assert len(matched["allocations"]) == 2

    other_expense = client.post(
        "/api/expenses/drafts",
        headers=dandi,
        json={
            "project_name": "AI 工具订阅",
            "actual_amount": 120,
            "expense_month": "2026-05",
            "category": "AI 项目",
        },
    ).json()
    duplicate_match = client.post(
        "/api/expense-allocations",
        headers=dandi,
        json={
            "expense_id": other_expense["id"],
            "attachment_id": first_attachment_id,
            "invoice_item_index": 0,
            "note": "",
        },
    )
    assert duplicate_match.status_code == 400
    assert "已经匹配" in duplicate_match.json()["detail"]

    pool = client.get("/api/invoice-pool", headers=dandi)
    assert pool.status_code == 200
    by_attachment = {item["attachment_id"]: item for item in pool.json()}
    assert by_attachment[first_attachment_id]["allocated_amount"] == 120
    assert by_attachment[first_attachment_id]["remaining_amount"] == 0
    assert by_attachment[second_attachment_id]["allocated_amount"] == 300
    assert by_attachment[second_attachment_id]["remaining_amount"] == 0


def test_cross_entity_invoice_buyer_requires_confirmation(tmp_path, monkeypatch):
    client = make_client(tmp_path, monkeypatch)
    dandi = auth_headers(client, "Dandi", "dandi123")

    attachment_id = insert_ocr_attachment(
        client,
        user_id=2,
        invoice_items=[
            {
                "buyer": "山途远智（上海）企业服务有限公司",
                "amount": 100,
                "invoice_number": "ENTITY-2",
                "date": "2026-05-20",
                "sub_type_description": "电子普通发票",
            }
        ],
        filename="ENTITY-2.pdf",
    )
    draft = client.post(
        "/api/expenses/drafts",
        headers=dandi,
        json={
            "project_name": "企业服务资料",
            "actual_amount": 100,
            "expense_month": "2026-05",
            "category": "办公采购",
        },
    )
    assert draft.status_code == 200

    rejected = client.post(
        "/api/expense-allocations",
        headers=dandi,
        json={
            "expense_id": draft.json()["id"],
            "attachment_id": attachment_id,
            "invoice_item_index": 0,
            "note": "",
        },
    )
    assert rejected.status_code == 400
    assert "公司主体" in rejected.json()["detail"]

    accepted = client.post(
        "/api/expense-allocations",
        headers=dandi,
        json={
            "expense_id": draft.json()["id"],
            "attachment_id": attachment_id,
            "invoice_item_index": 0,
            "note": "",
            "buyer_confirmed": True,
        },
    )
    assert accepted.status_code == 200
    assert accepted.json()["status"] == "pending"
    assert accepted.json()["invoice_buyer"] == "山途远智（上海）企业服务有限公司"


def test_partial_invoice_buyer_match_requires_manual_confirmation(tmp_path, monkeypatch):
    client = make_client(tmp_path, monkeypatch)
    dandi = auth_headers(client, "Dandi", "dandi123")
    attachment_id = insert_ocr_attachment(
        client,
        user_id=2,
        invoice_items=[
            {
                "buyer": "山途远智",
                "amount": 100,
                "invoice_number": "PARTIAL-BUYER",
                "date": "2026-05-21",
                "sub_type_description": "电子普通发票",
            }
        ],
        filename="partial-buyer.pdf",
    )
    draft = client.post(
        "/api/expenses/drafts",
        headers=dandi,
        json={
            "project_name": "客户资料打印",
            "actual_amount": 100,
            "expense_month": "2026-05",
            "category": "办公采购",
        },
    )
    assert draft.status_code == 200

    without_confirmation = client.post(
        "/api/expense-allocations",
        headers=dandi,
        json={
            "expense_id": draft.json()["id"],
            "attachment_id": attachment_id,
            "invoice_item_index": 0,
            "note": "",
        },
    )
    assert without_confirmation.status_code == 400
    assert "人工确认" in without_confirmation.json()["detail"]

    with_confirmation = client.post(
        "/api/expense-allocations",
        headers=dandi,
        json={
            "expense_id": draft.json()["id"],
            "attachment_id": attachment_id,
            "invoice_item_index": 0,
            "note": "",
            "buyer_confirmed": True,
        },
    )
    assert with_confirmation.status_code == 200
    assert with_confirmation.json()["status"] == "pending"


def test_submitted_expense_locks_invoices_but_allows_evidence(tmp_path, monkeypatch):
    client = make_client(tmp_path, monkeypatch)
    dandi = auth_headers(client, "Dandi", "dandi123")
    submitted = create_invoice_backed_expense(
        client,
        dandi,
        user_id=2,
        amount=100,
        project_name="客户拜访打车",
        category="差旅交通",
        invoice_number="LOCKED-1",
    )
    submit = client.post(f"/api/expenses/{submitted['id']}/submit", headers=dandi)
    assert submit.status_code == 200
    submitted = submit.json()
    extra_invoice_id = insert_ocr_attachment(
        client,
        user_id=2,
        invoice_items=[
            {
                "buyer": "上海山途远智信息科技有限公司",
                "amount": 20,
                "invoice_number": "LOCKED-2",
                "date": "2026-05-21",
                "sub_type_description": "电子普通发票",
            }
        ],
        filename="locked-extra.pdf",
    )

    append_invoice = client.post(
        "/api/expense-allocations",
        headers=dandi,
        json={
            "expense_id": submitted["id"],
            "attachment_id": extra_invoice_id,
            "invoice_item_index": 0,
            "note": "补充发票",
        },
    )
    assert append_invoice.status_code == 400
    assert "已提交" in append_invoice.json()["detail"]

    payment = upload_evidence_file(client, dandi, "late-payment.png", b"payment-image", "image/png")
    append_attachment = client.post(
        f"/api/expenses/{submitted['id']}/attachments",
        headers=dandi,
        json={"attachment_ids": [payment.json()[0]["id"]]},
    )
    assert append_attachment.status_code == 200
    assert {item["id"] for item in append_attachment.json()["attachments"]} == {payment.json()[0]["id"]}


def test_expense_keeps_transaction_attachments_out_of_invoice_pool(tmp_path, monkeypatch):
    client = make_client(tmp_path, monkeypatch)
    dandi = auth_headers(client, "Dandi", "dandi123")

    image_upload = upload_evidence_file(client, dandi, "payment.png", b"payment-image", "image/png")
    assert image_upload.status_code == 200
    image_id = image_upload.json()[0]["id"]
    content = client.get(f"/api/attachments/{image_id}/content", headers=dandi)
    assert content.status_code == 200
    assert content.content == b"payment-image"

    invoice_like_attachment_id = insert_ocr_attachment(
        client,
        user_id=2,
        invoice_items=[
            {
                "buyer": "上海山途远智信息科技有限公司",
                "amount": 199,
                "invoice_number": "PAYMENT-LIKE",
                "date": "2026-05-22",
                "sub_type_description": "交易截图识别结果",
            }
        ],
        filename="payment-like.png",
    )

    draft = client.post(
        "/api/expenses/drafts",
        headers=dandi,
        json={
            "project_name": "客户现场停车费",
            "actual_amount": 199,
            "expense_month": "2026-05",
            "category": "差旅交通",
        },
    )
    assert draft.status_code == 200

    link = client.post(
        f"/api/expenses/{draft.json()['id']}/attachments",
        headers=dandi,
        json={"attachment_ids": [image_id, invoice_like_attachment_id]},
    )
    assert link.status_code == 200
    linked = link.json()
    assert linked["status"] == "pending"
    assert {item["id"] for item in linked["attachments"]} == {image_id, invoice_like_attachment_id}

    pool = client.get("/api/invoice-pool", headers=dandi)
    assert pool.status_code == 200
    assert all(item["attachment_id"] != invoice_like_attachment_id for item in pool.json())


def test_employee_can_delete_transaction_attachment_from_expense(tmp_path, monkeypatch):
    client = make_client(tmp_path, monkeypatch)
    dandi = auth_headers(client, "Dandi", "dandi123")

    image_upload = upload_evidence_file(client, dandi, "wrong-payment.png", b"wrong-payment", "image/png")
    assert image_upload.status_code == 200
    draft = client.post(
        "/api/expenses/drafts",
        headers=dandi,
        json={
            "project_name": "客户现场停车费",
            "actual_amount": 199,
            "expense_month": "2026-05",
            "category": "差旅交通",
        },
    )
    assert draft.status_code == 200
    link = client.post(
        f"/api/expenses/{draft.json()['id']}/attachments",
        headers=dandi,
        json={"attachment_ids": [image_upload.json()[0]["id"]]},
    )
    assert link.status_code == 200
    assert len(link.json()["attachments"]) == 1

    delete = client.delete(
        f"/api/expenses/{draft.json()['id']}/attachments/{image_upload.json()[0]['id']}",
        headers=dandi,
    )
    assert delete.status_code == 200
    assert delete.json()["attachments"] == []
    missing = client.get(f"/api/attachments/{image_upload.json()[0]['id']}/content", headers=dandi)
    assert missing.status_code == 404


def test_employee_can_delete_linked_invoice_from_pending_expense(tmp_path, monkeypatch):
    client = make_client(tmp_path, monkeypatch)
    dandi = auth_headers(client, "Dandi", "dandi123")
    attachment_id = insert_ocr_attachment(
        client,
        user_id=2,
        invoice_items=[
            {
                "buyer": "上海山途远智信息科技有限公司",
                "amount": 120,
                "invoice_number": "DEL-INV-1",
                "date": "2026-05-23",
                "sub_type_description": "电子普通发票",
            }
        ],
        filename="delete-invoice.pdf",
    )
    draft = client.post(
        "/api/expenses/drafts",
        headers=dandi,
        json={
            "project_name": "客户拜访差旅",
            "actual_amount": 120,
            "expense_month": "2026-05",
            "category": "差旅交通",
        },
    )
    assert draft.status_code == 200
    expense_id = draft.json()["id"]
    match = client.post(
        "/api/expense-allocations",
        headers=dandi,
        json={
            "expense_id": expense_id,
            "attachment_id": attachment_id,
            "invoice_item_index": 0,
            "note": "",
        },
    )
    assert match.status_code == 200
    assert match.json()["allocation_count"] == 1

    delete = client.delete(f"/api/expenses/{expense_id}/invoice-attachments/{attachment_id}", headers=dandi)
    assert delete.status_code == 200
    assert delete.json()["allocation_count"] == 0
    assert delete.json()["invoice_number"] == ""
    missing = client.get(f"/api/attachments/{attachment_id}/content", headers=dandi)
    assert missing.status_code == 404


def test_employee_can_delete_unused_invoice_attachment(tmp_path, monkeypatch):
    client = make_client(tmp_path, monkeypatch)
    dandi = auth_headers(client, "Dandi", "dandi123")
    attachment_id = insert_ocr_attachment(
        client,
        user_id=2,
        invoice_items=[
            {
                "buyer": "上海山途远智信息科技有限公司",
                "amount": 88,
                "invoice_number": "WRONG-1",
                "date": "2026-05-23",
                "sub_type_description": "电子普通发票",
            }
        ],
        filename="wrong-invoice.pdf",
    )
    pool = client.get("/api/invoice-pool", headers=dandi)
    assert any(item["attachment_id"] == attachment_id for item in pool.json())

    delete = client.delete(f"/api/attachments/{attachment_id}", headers=dandi)
    assert delete.status_code == 200
    assert delete.json() == {"deleted": True}
    pool_after = client.get("/api/invoice-pool", headers=dandi)
    assert all(item["attachment_id"] != attachment_id for item in pool_after.json())
    missing = client.get(f"/api/attachments/{attachment_id}/content", headers=dandi)
    assert missing.status_code == 404


def test_deleting_draft_expense_releases_matched_invoice(tmp_path, monkeypatch):
    client = make_client(tmp_path, monkeypatch)
    dandi = auth_headers(client, "Dandi", "dandi123")
    attachment_id = insert_ocr_attachment(
        client,
        user_id=2,
        invoice_items=[
            {
                "buyer": "上海山途远智信息科技有限公司",
                "amount": 120,
                "invoice_number": "PARTIAL-1",
                "date": "2026-05-23",
                "sub_type_description": "电子普通发票",
            }
        ],
        filename="partial-invoice.pdf",
    )
    draft = client.post(
        "/api/expenses/drafts",
        headers=dandi,
        json={
            "project_name": "客户拜访差旅",
            "actual_amount": 300,
            "expense_month": "2026-05",
            "category": "差旅交通",
        },
    )
    assert draft.status_code == 200
    match = client.post(
        "/api/expense-allocations",
        headers=dandi,
        json={
            "expense_id": draft.json()["id"],
            "attachment_id": attachment_id,
            "invoice_item_index": 0,
            "note": "",
        },
    )
    assert match.status_code == 200
    assert match.json()["status"] == "pending"

    used_pool = client.get("/api/invoice-pool", headers=dandi)
    used_item = next(item for item in used_pool.json() if item["attachment_id"] == attachment_id)
    assert used_item["remaining_amount"] == 0

    delete = client.delete(f"/api/expenses/{draft.json()['id']}", headers=dandi)
    assert delete.status_code == 200
    assert delete.json() == {"deleted": True}
    expenses = client.get("/api/expenses", headers=dandi)
    assert all(item["id"] != draft.json()["id"] for item in expenses.json())

    released_pool = client.get("/api/invoice-pool", headers=dandi)
    released_item = next(item for item in released_pool.json() if item["attachment_id"] == attachment_id)
    assert released_item["remaining_amount"] == 120


def test_invoice_match_requires_reason_when_invoice_total_exceeds_expense(tmp_path, monkeypatch):
    client = make_client(tmp_path, monkeypatch)
    dandi = auth_headers(client, "Dandi", "dandi123")
    attachment_id = insert_ocr_attachment(
        client,
        user_id=2,
        invoice_items=[
            {
                "buyer": "上海山途远智信息科技有限公司",
                "amount": 120,
                "invoice_number": "OVER-1",
                "date": "2026-05-23",
                "sub_type_description": "电子普通发票",
            }
        ],
        filename="over-invoice.pdf",
    )

    draft = client.post(
        "/api/expenses/drafts",
        headers=dandi,
        json={
            "project_name": "市场物料垫付",
            "actual_amount": 100,
            "expense_month": "2026-05",
            "category": "市场活动",
        },
    ).json()

    no_reason = client.post(
        "/api/expense-allocations/batch",
        headers=dandi,
        json={
            "expense_id": draft["id"],
            "invoices": [{"attachment_id": attachment_id, "invoice_item_index": 0}],
            "note": "",
        },
    )
    assert no_reason.status_code == 400
    assert "说明" in no_reason.json()["detail"]

    with_reason = client.post(
        "/api/expense-allocations/batch",
        headers=dandi,
        json={
            "expense_id": draft["id"],
            "invoices": [{"attachment_id": attachment_id, "invoice_item_index": 0}],
            "note": "用同项目大额发票替票",
        },
    )
    assert with_reason.status_code == 200
    body = with_reason.json()
    assert body["status"] == "pending"
    assert body["is_substitute"] is True
    assert body["substitute_reason"] == "用同项目大额发票替票"


def test_invoice_allocation_requires_reason_when_amount_exceeds_expense(tmp_path, monkeypatch):
    client = make_client(tmp_path, monkeypatch)
    dandi = auth_headers(client, "Dandi", "dandi123")
    attachment_id = insert_ocr_attachment(
        client,
        user_id=2,
        invoice_items=[{"buyer": "上海山途远智信息科技有限公司", "amount": 120}],
        filename="amount-mismatch.pdf",
    )

    draft = client.post(
        "/api/expenses/drafts",
        headers=dandi,
        json={
            "project_name": "市场物料垫付",
            "actual_amount": 100,
            "expense_month": "2026-05",
            "category": "市场活动",
        },
    )
    assert draft.status_code == 200

    response = client.post(
        "/api/expense-allocations",
        headers=dandi,
        json={
            "expense_id": draft.json()["id"],
            "attachment_id": attachment_id,
            "invoice_item_index": 0,
            "note": "",
        },
    )
    assert response.status_code == 400
    assert "说明" in response.json()["detail"]


def test_single_invoice_upload_rejects_when_ocr_not_configured(tmp_path, monkeypatch):
    client = make_client(tmp_path, monkeypatch)
    dandi = auth_headers(client, "Dandi", "dandi123")

    response = upload_file(client, dandi, "invoice.pdf", b"fake-invoice", "application/pdf")
    assert response.status_code == 400
    assert response.json()["detail"] == "腾讯云 OCR 未配置，无法识别发票"


def test_batch_upload_accepts_multiple_files(tmp_path, monkeypatch):
    client = make_client(tmp_path, monkeypatch)
    dandi = auth_headers(client, "Dandi", "dandi123")

    response = client.post(
        "/api/attachments/batch",
        headers=dandi,
        files=[
            ("files", ("invoice-a.pdf", b"invoice-a", "application/pdf")),
            ("files", ("invoice-b.pdf", b"invoice-b", "application/pdf")),
        ],
    )

    assert response.status_code == 200
    body = response.json()
    assert len(body) == 2
    assert body[0]["ocr_status"] == "skipped"
    assert body[1]["ocr_status"] == "skipped"
    assert body[0]["pool_status"] == "staged"
    assert body[1]["pool_status"] == "staged"


def test_staged_invoice_enters_pool_only_after_button_action(tmp_path, monkeypatch):
    client = make_client(tmp_path, monkeypatch)
    dandi = auth_headers(client, "Dandi", "dandi123")
    attachment_id = insert_ocr_attachment(
        client,
        user_id=2,
        invoice_items=[
            {
                "buyer": "上海山途远智信息科技有限公司",
                "amount": 88.6,
                "invoice_number": "POOL-1",
                "date": "2026-05-18",
                "sub_type_description": "电子普通发票",
            }
        ],
        filename="staged-invoice.pdf",
        pool_status="staged",
    )

    pool = client.get("/api/invoice-pool", headers=dandi)
    assert pool.status_code == 200
    assert all(item["attachment_id"] != attachment_id for item in pool.json())

    publish = client.post("/api/attachments/pool", headers=dandi, json={"attachment_ids": [attachment_id]})
    assert publish.status_code == 200
    assert publish.json()[0]["pool_status"] == "pooled"

    pool_after = client.get("/api/invoice-pool", headers=dandi)
    assert any(item["attachment_id"] == attachment_id for item in pool_after.json())


def test_staged_invoice_can_bind_directly_without_entering_pool(tmp_path, monkeypatch):
    client = make_client(tmp_path, monkeypatch)
    dandi = auth_headers(client, "Dandi", "dandi123")
    attachment_id = insert_ocr_attachment(
        client,
        user_id=2,
        invoice_items=[
            {
                "buyer": "上海山途远智信息科技有限公司",
                "amount": 88.6,
                "invoice_number": "DIRECT-1",
                "date": "2026-05-18",
                "sub_type_description": "电子普通发票",
            }
        ],
        filename="direct-staged-invoice.pdf",
        pool_status="staged",
    )

    draft = client.post(
        "/api/expenses/drafts",
        headers=dandi,
        json={
            "project_name": "客户拜访打车",
            "actual_amount": 88.6,
            "expense_month": "2026-05",
            "category": "差旅交通",
        },
    )
    assert draft.status_code == 200

    match = client.post(
        "/api/expense-allocations/batch",
        headers=dandi,
        json={
            "expense_id": draft.json()["id"],
            "invoices": [{"attachment_id": attachment_id, "invoice_item_index": 0}],
            "note": "",
        },
    )
    assert match.status_code == 200
    assert match.json()["status"] == "pending"

    pool = client.get("/api/invoice-pool", headers=dandi)
    assert all(item["attachment_id"] != attachment_id for item in pool.json())


def test_batch_expense_creates_one_record_per_invoice_item(tmp_path, monkeypatch):
    client = make_client(tmp_path, monkeypatch)
    dandi = auth_headers(client, "Dandi", "dandi123")
    attachment_id = insert_ocr_attachment(
        client,
        user_id=2,
        invoice_items=[
            {
                "buyer": "上海山途远智信息科技有限公司",
                "amount": 88.6,
                "invoice_number": "NO-1",
                "date": "2026-05-18",
                "sub_type_description": "电子普通发票",
            },
            {
                "buyer": "上海山途远智信息科技有限公司",
                "amount": 12.3,
                "invoice_number": "NO-2",
                "date": "2026-05-18",
                "sub_type_description": "出租车票",
            },
        ],
    )

    response = client.post(
        "/api/expenses/batch",
        headers=dandi,
        json={
            "items": [
                {
                    "attachment_id": attachment_id,
                    "invoice_item_index": 0,
                    "category": "办公采购",
                    "expense_month": "2026-05",
                    "actual_amount": None,
                    "is_substitute": False,
                    "substitute_reason": "",
                    "note": "",
                },
                {
                    "attachment_id": attachment_id,
                    "invoice_item_index": 1,
                    "category": "差旅交通",
                    "expense_month": "2026-05",
                    "actual_amount": None,
                    "is_substitute": False,
                    "substitute_reason": "",
                    "note": "",
                },
            ]
        },
    )

    assert response.status_code == 200
    body = response.json()
    assert len(body) == 2
    assert body[0]["invoice_amount"] == 88.6
    assert body[0]["actual_amount"] == 88.6
    assert body[1]["invoice_number"] == "NO-2"


def test_batch_expense_rejects_wrong_company_title(tmp_path, monkeypatch):
    client = make_client(tmp_path, monkeypatch)
    dandi = auth_headers(client, "Dandi", "dandi123")
    attachment_id = insert_ocr_attachment(
        client,
        user_id=2,
        invoice_items=[{"buyer": "别的公司有限公司", "amount": 88.6}],
        filename="wrong-title.pdf",
    )

    response = client.post(
        "/api/expenses/batch",
        headers=dandi,
        json={
            "items": [
                {
                    "attachment_id": attachment_id,
                    "invoice_item_index": 0,
                    "category": "AI 项目",
                    "expense_month": "2026-05",
                    "actual_amount": None,
                    "is_substitute": False,
                    "substitute_reason": "",
                    "note": "",
                }
            ]
        },
    )

    assert response.status_code == 400
    assert "抬头" in response.json()["detail"]


def test_admin_can_create_update_and_deactivate_user(tmp_path, monkeypatch):
    client = make_client(tmp_path, monkeypatch)
    admin = auth_headers(client, "admin", "admin123")

    create = client.post(
        "/api/admin/users",
        headers=admin,
        json={
            "username": "carol",
            "email": "Carol@Mentitrek.com ",
            "password": "carol123",
            "role": "employee",
            "employee_name": "Carol Wang",
            "company_entity": "上海山途远智信息科技有限公司",
        },
    )
    assert create.status_code == 200
    assert create.json()["email"] == "carol@mentitrek.com"
    user_id = create.json()["id"]

    update = client.patch(
        f"/api/admin/users/{user_id}",
        headers=admin,
        json={
            "email": "carol.wang@mentitrek.com",
            "company_entity": "山途远智（上海）企业服务有限公司",
            "password": "newpass123",
        },
    )
    assert update.status_code == 200
    assert update.json()["email"] == "carol.wang@mentitrek.com"
    assert update.json()["company_entity"] == "山途远智（上海）企业服务有限公司"

    carol = auth_headers(client, "carol", "newpass123")
    assert client.get("/api/me", headers=carol).status_code == 200

    deactivate = client.delete(f"/api/admin/users/{user_id}", headers=admin)
    assert deactivate.status_code == 200
    assert deactivate.json()["is_active"] is False
    disabled_login = client.post("/api/auth/login", json={"username": "carol", "password": "newpass123"})
    assert disabled_login.status_code == 403


def test_sso_admin_can_preprovision_employee_without_local_password(tmp_path, monkeypatch):
    client = make_client(tmp_path, monkeypatch, auth_mode="sso")
    from app.security import create_token

    client.cookies.set("jetbao_session", create_token(1, "test-secret"))
    create = client.post(
        "/api/admin/users",
        json={
            "username": "new-employee",
            "email": "new.employee@mentitrek.com",
            "role": "employee",
            "employee_name": "New Employee",
            "company_entity": "上海山途远智信息科技有限公司",
        },
    )

    assert create.status_code == 200
    assert create.json()["email"] == "new.employee@mentitrek.com"
    assert create.json()["identity_id"] is None


def test_admin_user_company_entity_must_be_allowed(tmp_path, monkeypatch):
    client = make_client(tmp_path, monkeypatch)
    admin = auth_headers(client, "admin", "admin123")

    create = client.post(
        "/api/admin/users",
        headers=admin,
        json={
            "username": "bad-company",
            "password": "bad123",
            "role": "employee",
            "employee_name": "Bad Company",
            "company_entity": "杭州示例信息有限公司",
        },
    )

    assert create.status_code == 400
    assert "公司主体" in create.json()["detail"]


def test_admin_can_preview_expense_attachments_for_review(tmp_path, monkeypatch):
    client = make_client(tmp_path, monkeypatch)
    dandi = auth_headers(client, "Dandi", "dandi123")
    admin = auth_headers(client, "admin", "admin123")

    payment = upload_evidence_file(client, dandi, "payment-proof.png", b"payment-proof", "image/png")
    assert payment.status_code == 200
    payment_id = payment.json()[0]["id"]

    expense = create_invoice_backed_expense(client, dandi, user_id=2, amount=300)
    expense_id = expense["id"]
    assert expense["status"] == "pending"

    link = client.post(
        f"/api/expenses/{expense_id}/attachments",
        headers=dandi,
        json={"attachment_ids": [payment_id]},
    )
    assert link.status_code == 200

    submit = client.post(f"/api/expenses/{expense_id}/submit", headers=dandi)
    assert submit.status_code == 200
    assert submit.json()["status"] == "matched"

    review = client.get(f"/api/admin/expenses/{expense_id}", headers=admin)
    assert review.status_code == 200
    body = review.json()
    assert body["status"] == "matched"
    assert len(body["attachments"]) == 1
    assert body["attachments"][0]["id"] == payment_id
    assert len(body["invoice_attachments"]) == 1
    invoice_id = body["invoice_attachments"][0]["id"]

    employee_content = client.get(f"/api/attachments/{payment_id}/content", headers=admin)
    assert employee_content.status_code == 200
    assert employee_content.content == b"payment-proof"

    invoice_content = client.get(f"/api/attachments/{invoice_id}/content", headers=admin)
    assert invoice_content.status_code == 200

    created_employee = client.post(
        "/api/admin/users",
        headers=admin,
        json={
            "username": "employee-review",
            "password": "employee-pass",
            "role": "employee",
            "employee_name": "普通员工",
            "company_entity": "上海山途远智信息科技有限公司",
        },
    )
    assert created_employee.status_code == 200
    employee = auth_headers(client, "employee-review", "employee-pass")
    forbidden = client.get(f"/api/admin/expenses/{expense_id}", headers=employee)
    assert forbidden.status_code == 403
