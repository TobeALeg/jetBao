from __future__ import annotations

import json
import re
import sqlite3
import zipfile
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.expense_month_filter import apply_expense_month_filter, expense_period_label, normalize_month, normalize_month_part, normalize_year


INVALID_PATH_CHARS = re.compile(r'[\\/:*?"<>|\r\n]+')
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp", ".heic"}


@dataclass(frozen=True)
class ExportDocument:
    source_path: Path
    package_path: str
    original_filename: str
    exists: bool
    amount_hint: float | None = None
    attachment_kind: str = "非发票附件"


@dataclass(frozen=True)
class ExportInvoice:
    allocation: sqlite3.Row
    invoice_item: dict[str, Any]
    document: ExportDocument
    project_name: str = ""
    expense_status: str = ""


PENDING_APPROVAL_NOTE = "等待通过"


@dataclass(frozen=True)
class ExportBundle:
    expense: sqlite3.Row
    group_id: str
    expense_name: str
    folder_path: str
    invoices: list[ExportInvoice]
    transaction_attachments: list[ExportDocument]
    member_expenses: tuple[sqlite3.Row, ...] = ()


def build_export_package(
    connection: sqlite3.Connection,
    month: str | None,
    company_entity: str | None,
    year: str | None = None,
    month_part: str | None = None,
) -> tuple[BytesIO, str]:
    bundles = load_export_bundles(connection, month, company_entity, year=year, month_part=month_part)
    workbook = build_export_workbook(bundles, month, company_entity, year=year, month_part=month_part)

    month_name = export_period_label(month, year, month_part)
    root_folder = f"山途远智{month_name}报销"
    output = BytesIO()
    written_files: set[str] = set()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(f"{month_name}报销明细.xlsx", workbook.getvalue())
        for bundle in bundles:
            for document in [*bundle.transaction_attachments, *[invoice.document for invoice in bundle.invoices]]:
                if not document.exists or document.package_path in written_files:
                    continue
                archive.write(document.source_path, document.package_path)
                written_files.add(document.package_path)

    output.seek(0)
    return output, f"{root_folder}明细包.zip"


def build_export_workbook(
    bundles: list[ExportBundle],
    month: str | None,
    company_entity: str | None,
    year: str | None = None,
    month_part: str | None = None,
) -> BytesIO:
    workbook = Workbook()
    overview = workbook.active
    overview.title = "总览"
    summary = workbook.create_sheet("报销项汇总")
    invoices_sheet = workbook.create_sheet("发票明细")
    attachments_sheet = workbook.create_sheet("附件与待核对")

    period_label = export_period_label(month, year, month_part)
    _fill_overview(overview, workbook, bundles, period_label, company_entity)
    _fill_summary(summary, bundles)
    _fill_invoices(invoices_sheet, bundles)
    _fill_attachments(attachments_sheet, bundles)

    for sheet in workbook.worksheets:
        if sheet.title.startswith("_"):
            continue
        _style_sheet(sheet)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def load_export_bundles(
    connection: sqlite3.Connection,
    month: str | None,
    company_entity: str | None,
    year: str | None = None,
    month_part: str | None = None,
) -> list[ExportBundle]:
    expenses = _fetch_exportable_expenses(connection, month, company_entity, year=year, month_part=month_part)
    month_name = export_period_label(month, year, month_part)
    folder_month = folder_month_label(month, year, month_part)
    root_folder = f"山途远智{month_name}报销"
    used_paths: set[str] = set()
    bundles: list[ExportBundle] = []

    # Group by person + reimbursement category (差旅交通 / AI 项目 / …).
    grouped: dict[tuple[str, str], list[sqlite3.Row]] = {}
    for expense in expenses:
        employee = str(expense["employee_name"] or "").strip() or "未命名员工"
        category = str(expense["category"] or "").strip() or "未分类"
        grouped.setdefault((employee, category), []).append(expense)

    for index, ((employee, category), members) in enumerate(grouped.items(), start=1):
        group_id = f"G{index:02d}"
        representative = members[0]
        category_name = category
        category_folder = f"{safe_path_part(category, '未分类')}{folder_month}报销"
        employee_folder = f"{safe_path_part(employee, '未命名员工')}{folder_month}报销"
        folder_path = f"{root_folder}/{category_folder}/{employee_folder}"

        transactions: list[ExportDocument] = []
        invoices: list[ExportInvoice] = []
        for expense in members:
            item_name = _expense_name(expense)
            for row in _fetch_transaction_attachments(connection, expense["id"]):
                transactions.append(
                    _document_named_by_item(
                        row,
                        folder_path,
                        item_name,
                        used_paths,
                        name_role="佐证材料",
                        amount_hint=float(expense["actual_amount"] or 0),
                        attachment_kind=_attachment_kind(row["original_filename"]),
                    )
                )
            for allocation in _fetch_invoice_allocations(connection, expense["id"]):
                document = _document_named_by_item(
                    allocation,
                    folder_path,
                    item_name,
                    used_paths,
                    name_role="发票",
                    amount_hint=float(expense["actual_amount"] or 0),
                    attachment_kind="发票",
                )
                invoices.append(
                    ExportInvoice(
                        allocation=allocation,
                        invoice_item=_invoice_item_from_allocation(allocation),
                        document=document,
                        project_name=item_name,
                        expense_status=str(expense["status"] or ""),
                    )
                )

        bundles.append(
            ExportBundle(
                expense=representative,
                group_id=group_id,
                expense_name=category_name,
                folder_path=folder_path,
                invoices=invoices,
                transaction_attachments=transactions,
                member_expenses=tuple(members),
            )
        )
    return bundles


def month_label(month: str | None) -> str:
    if month and re.fullmatch(r"\d{4}-\d{2}", month):
        return f"{int(month[-2:])}月"
    return "全部"


def export_period_label(
    month: str | None,
    year: str | None = None,
    month_part: str | None = None,
) -> str:
    period = expense_period_label(
        normalize_month(month),
        normalize_year(year),
        normalize_month_part(month_part),
    )
    if re.fullmatch(r"\d{4}-\d{2}", period):
        return f"{int(period[-2:])}月"
    return period


def folder_month_label(
    month: str | None,
    year: str | None = None,
    month_part: str | None = None,
) -> str:
    """Zero-padded month tag used in nested ZIP folders, e.g. 04月."""
    period = expense_period_label(
        normalize_month(month),
        normalize_year(year),
        normalize_month_part(month_part),
    )
    if re.fullmatch(r"\d{4}-\d{2}", period):
        return f"{period[-2:]}月"
    if month_part and re.fullmatch(r"\d{1,2}", month_part):
        return f"{int(month_part):02d}月"
    return export_period_label(month, year, month_part)


def company_alias(company_entity: str) -> str:
    for keyword in ("信息科技", "企业服务", "企业咨询"):
        if keyword in company_entity:
            return keyword
    cleaned = re.sub(r"（.*?）|\(.*?\)|有限公司|有限责任公司|上海|山途远智", "", company_entity)
    return safe_path_part(cleaned, "公司")


def safe_path_part(value: str | None, fallback: str) -> str:
    cleaned = INVALID_PATH_CHARS.sub("_", (value or "").strip())
    cleaned = re.sub(r"\s+", " ", cleaned).strip(" ._")
    return cleaned[:80] or fallback


def format_amount(value: float | int | str | None) -> str:
    number = round(float(value or 0), 2)
    if number.is_integer():
        return str(int(number))
    return f"{number:.2f}".rstrip("0").rstrip(".")


def _fetch_exportable_expenses(
    connection: sqlite3.Connection,
    month: str | None,
    company_entity: str | None,
    year: str | None = None,
    month_part: str | None = None,
) -> list[sqlite3.Row]:
    where = ["expenses.status IN ('matched', 'reviewed')"]
    params: list[object] = []
    apply_expense_month_filter(
        where,
        params,
        month=normalize_month(month),
        year=normalize_year(year),
        month_part=normalize_month_part(month_part),
    )
    if company_entity:
        where.append("expenses.company_entity = ?")
        params.append(company_entity)

    return connection.execute(
        f"""
        SELECT
            expenses.*,
            users.employee_name,
            users.company_entity AS user_company_entity
        FROM expenses
        JOIN users ON users.id = expenses.user_id
        WHERE {" AND ".join(where)}
        ORDER BY
            users.employee_name COLLATE NOCASE,
            expenses.category COLLATE NOCASE,
            expenses.created_at,
            expenses.id
        """,
        params,
    ).fetchall()


def _fetch_transaction_attachments(connection: sqlite3.Connection, expense_id: int) -> list[sqlite3.Row]:
    return connection.execute(
        """
        SELECT attachments.*
        FROM attachments
        JOIN expense_attachments ON expense_attachments.attachment_id = attachments.id
        WHERE expense_attachments.expense_id = ?
        UNION
        SELECT * FROM attachments WHERE expense_id = ?
        ORDER BY created_at, id
        """,
        (expense_id, expense_id),
    ).fetchall()


def _fetch_invoice_allocations(connection: sqlite3.Connection, expense_id: int) -> list[sqlite3.Row]:
    return connection.execute(
        """
        SELECT
            expense_invoice_allocations.*,
            attachments.original_filename,
            attachments.stored_path,
            attachments.ocr_result
        FROM expense_invoice_allocations
        JOIN attachments ON attachments.id = expense_invoice_allocations.attachment_id
        WHERE expense_invoice_allocations.expense_id = ?
        ORDER BY expense_invoice_allocations.id
        """,
        (expense_id,),
    ).fetchall()


def _invoice_item_from_allocation(allocation: sqlite3.Row) -> dict[str, Any]:
    try:
        result = json.loads(allocation["ocr_result"] or "{}")
    except json.JSONDecodeError:
        return {}
    items = result.get("invoice_items")
    index = int(allocation["invoice_item_index"])
    if not isinstance(items, list) or index >= len(items) or not isinstance(items[index], dict):
        return {}
    return items[index]


def _attachment_kind(filename: str | None) -> str:
    suffix = Path(filename or "").suffix.lower()
    if suffix in IMAGE_EXTENSIONS:
        return "图片/截图附件"
    return "非发票附件"


def _document_named_by_item(
    attachment: sqlite3.Row,
    folder: str,
    item_name: str,
    used_paths: set[str],
    *,
    name_role: str,
    amount_hint: float | None = None,
    attachment_kind: str = "非发票附件",
) -> ExportDocument:
    return ExportDocument(
        source_path=Path(attachment["stored_path"]),
        package_path=_package_item_path(
            folder,
            item_name,
            name_role,
            attachment["original_filename"],
            used_paths,
            amount=amount_hint,
        ),
        original_filename=attachment["original_filename"],
        exists=Path(attachment["stored_path"]).exists(),
        amount_hint=amount_hint,
        attachment_kind=attachment_kind,
    )


def _package_item_path(
    folder: str,
    item_name: str,
    name_role: str,
    original_filename: str,
    used_paths: set[str],
    *,
    amount: float | None = None,
) -> str:
    """Build ZIP filename like「AI项目佐证材料700元.png」or「AI项目发票700元.pdf」."""
    suffix = Path(original_filename or "").suffix
    amount_tag = f"{format_amount(amount)}元" if amount is not None and float(amount or 0) > 0 else ""
    base = safe_path_part(f"{item_name}{name_role}{amount_tag}", f"报销{name_role}{amount_tag}")
    path = f"{folder}/{base}{suffix}".replace("//", "/")
    if path not in used_paths:
        used_paths.add(path)
        return path

    counter = 2
    while True:
        candidate = f"{folder}/{base} ({counter}){suffix}".replace("//", "/")
        if candidate not in used_paths:
            used_paths.add(candidate)
            return candidate
        counter += 1


def _expense_name(expense: sqlite3.Row) -> str:
    return (
        str(expense["project_name"] or "").strip()
        or str(expense["note"] or "").strip()
        or str(expense["category"] or "").strip()
        or "报销项"
    )


def _invoice_text(item: dict[str, Any], *keys: str) -> str:
    for key in keys:
        value = item.get(key)
        if value is not None and str(value).strip():
            return str(value).strip()
    return ""


def _invoice_total(bundle: ExportBundle) -> float:
    return round(sum(float(invoice.allocation["invoice_amount"]) for invoice in bundle.invoices), 2)


def _attachment_count(bundle: ExportBundle) -> int:
    invoice_files = {invoice.document.package_path for invoice in bundle.invoices}
    return len(bundle.transaction_attachments) + len(invoice_files)


def _fill_overview(
    sheet,
    workbook: Workbook,
    bundles: list[ExportBundle],
    period_label: str,
    company_entity: str | None,
) -> None:
    metrics = _overview_metrics(bundles)
    sheet.merge_cells("A1:G2")
    sheet["A1"] = f"{period_label}发票/报销整理总览"
    sheet["A4"] = "本次报销金额合计"
    sheet["A5"] = metrics["actual_total"]
    sheet["C4"] = "票面金额合计"
    sheet["C5"] = metrics["invoice_total"]
    sheet["E4"] = "涉及人员"
    sheet["E5"] = metrics["employee_count"]
    sheet["G4"] = "发票张数"
    sheet["G5"] = metrics["invoice_count"]

    company_start = 7
    company_rows = _overview_company_rows(bundles)
    _append_overview_table(
        sheet,
        company_start,
        ["公司主体", "人员数", "报销项数", "发票张数", "票面金额合计", "本次报销金额", "票面-报销差异"],
        company_rows,
    )

    detail_start = company_start + max(len(company_rows), 1) + 3
    employee_rows = _overview_employee_rows(bundles)
    _append_overview_table(
        sheet,
        detail_start,
        ["公司主体", "人员", "报销项数", "发票张数", "票面金额合计", "本次报销金额", "票面-报销差异"],
        employee_rows,
    )

    _add_overview_chart(sheet, workbook, employee_rows)

    note_start = detail_start + max(len(employee_rows), 1) + 3
    sheet.merge_cells(start_row=note_start, start_column=1, end_row=note_start, end_column=7)
    sheet.cell(row=note_start, column=1, value="核对提示")
    for index, note in enumerate(_overview_notes(bundles), start=1):
        row = note_start + index
        sheet.cell(row=row, column=1, value=index)
        sheet.cell(row=row, column=2, value=note)
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=7)


def _add_overview_chart(sheet, workbook: Workbook, employee_rows: list[list[object]]) -> None:
    if not employee_rows:
        return

    # WPS/部分 Excel 对「隐藏列引用」的图兼容较差，把源数据放到隐藏工作表里。
    data_sheet_name = "_chart_data"
    if data_sheet_name in workbook.sheetnames:
        data_sheet = workbook[data_sheet_name]
        workbook.remove(data_sheet)
    data_sheet = workbook.create_sheet(data_sheet_name)
    data_sheet["A1"] = "人员"
    data_sheet["B1"] = "票面金额合计"
    data_sheet["C1"] = "本次报销金额"
    for index, row_values in enumerate(employee_rows, start=2):
        data_sheet.cell(row=index, column=1, value=row_values[1])
        data_sheet.cell(row=index, column=2, value=float(row_values[4] or 0))
        data_sheet.cell(row=index, column=3, value=float(row_values[5] or 0))
    data_sheet.sheet_state = "hidden"

    chart = BarChart()
    chart.type = "col"
    chart.grouping = "clustered"
    chart.title = "按人员本次报销金额"
    chart.style = 10
    chart.legend.position = "b"
    data = Reference(data_sheet, min_col=2, min_row=1, max_col=3, max_row=1 + len(employee_rows))
    cats = Reference(data_sheet, min_col=1, min_row=2, max_row=1 + len(employee_rows))
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.width = 15
    chart.height = 10
    sheet.add_chart(chart, "I7")


def _append_overview_table(sheet, start_row: int, headers: list[str], rows: list[list[object]]) -> None:
    for column, header in enumerate(headers, start=1):
        sheet.cell(row=start_row, column=column, value=header)
    if not rows:
        sheet.cell(row=start_row + 1, column=1, value="无数据")
        sheet.merge_cells(start_row=start_row + 1, start_column=1, end_row=start_row + 1, end_column=len(headers))
        return
    for row_offset, row_values in enumerate(rows, start=1):
        for column, value in enumerate(row_values, start=1):
            sheet.cell(row=start_row + row_offset, column=column, value=value)


def _overview_metrics(bundles: list[ExportBundle]) -> dict[str, float | int]:
    actual_total = round(sum(_actual_total(bundle) for bundle in bundles), 2)
    invoice_total = round(sum(_invoice_total(bundle) for bundle in bundles), 2)
    return {
        "actual_total": actual_total,
        "invoice_total": invoice_total,
        "employee_count": len({(bundle.expense["company_entity"], bundle.expense["employee_name"]) for bundle in bundles}),
        "invoice_count": sum(len(bundle.invoices) for bundle in bundles),
    }


def _overview_company_rows(bundles: list[ExportBundle]) -> list[list[object]]:
    rows = []
    for company, group in _group_bundles(bundles, lambda bundle: bundle.expense["company_entity"]).items():
        rows.append(_overview_row(company_alias(company), group))
    return rows


def _overview_employee_rows(bundles: list[ExportBundle]) -> list[list[object]]:
    rows = []
    grouped = _group_bundles(
        bundles,
        lambda bundle: (bundle.expense["company_entity"], bundle.expense["employee_name"]),
    )
    for (company, employee_name), group in grouped.items():
        rows.append(_overview_row(company_alias(company), group, employee_name))
    return rows


def _overview_row(company_name: str, bundles: list[ExportBundle], employee_name: str | None = None) -> list[object]:
    actual_total = round(sum(_actual_total(bundle) for bundle in bundles), 2)
    invoice_total = round(sum(_invoice_total(bundle) for bundle in bundles), 2)
    return [
        company_name,
        len({bundle.expense["employee_name"] for bundle in bundles}) if employee_name is None else employee_name,
        len(bundles),
        sum(len(bundle.invoices) for bundle in bundles),
        invoice_total,
        actual_total,
        round(invoice_total - actual_total, 2),
    ]


def _group_bundles(bundles: list[ExportBundle], key_fn) -> dict:
    grouped = {}
    for bundle in bundles:
        key = key_fn(bundle)
        grouped.setdefault(key, []).append(bundle)
    return dict(sorted(grouped.items(), key=lambda item: str(item[0])))


def _actual_total(bundle: ExportBundle) -> float:
    members = bundle.member_expenses or (bundle.expense,)
    return round(sum(float(expense["actual_amount"] or 0) for expense in members), 2)


def _source_leaf_filename(document: ExportDocument) -> str:
    """Innermost filename of the uploaded source (e.g. DeepSeekAPI服务10元.jpg)."""
    if document.original_filename:
        return Path(document.original_filename).name
    return Path(document.source_path).name


def _overview_notes(bundles: list[ExportBundle]) -> list[str]:
    substitute_count = sum(1 for bundle in bundles if bundle.expense["is_substitute"])
    missing_documents = [
        document
        for bundle in bundles
        for document in [*bundle.transaction_attachments, *[invoice.document for invoice in bundle.invoices]]
        if not document.exists
    ]
    notes = [
        "“本次报销金额”按花费项目实际金额归集；替票场景不直接按发票票面金额报销。",
        "发票开票日期和报销归属按已提交记录整理，待补材料不会进入正式导出。",
        "“附件与待核对”列出交易记录、支付截图、详情截图等非正式发票材料。",
    ]
    if substitute_count:
        notes.append(f"共有 {substitute_count} 个替票/金额差异项，请在报销项汇总和发票明细中核对说明。")
    if missing_documents:
        notes.append(f"共有 {len(missing_documents)} 个源附件文件缺失，请在附件与待核对中补齐。")
    return notes


def _is_pending_approval(expense: sqlite3.Row | ExportInvoice | str) -> bool:
    if isinstance(expense, ExportInvoice):
        status = expense.expense_status
    elif isinstance(expense, str):
        status = expense
    else:
        status = str(expense["status"] or "")
    return status == "matched"


def _bundle_pending_approval(bundle: ExportBundle) -> bool:
    members = bundle.member_expenses or (bundle.expense,)
    return any(_is_pending_approval(member) for member in members)


def _compose_remark(existing: str, pending_approval: bool) -> str:
    note = (existing or "").strip()
    if not pending_approval:
        return note
    if PENDING_APPROVAL_NOTE in note:
        return note
    if note:
        return f"{note}；{PENDING_APPROVAL_NOTE}"
    return PENDING_APPROVAL_NOTE


def _fill_summary(sheet, bundles: list[ExportBundle]) -> None:
    sheet.append(
        [
            "组ID",
            "公司主体",
            "购买方名称",
            "人员",
            "报销项",
            "票据状态",
            "发票张数",
            "票面金额合计",
            "本次报销金额",
            "票面-报销差异",
            "附件数",
            "备注",
        ]
    )
    for bundle in bundles:
        invoice_total = _invoice_total(bundle)
        actual_amount = _actual_total(bundle)
        buyer = next(
            (
                str(expense["invoice_buyer"] or "").strip()
                for expense in (bundle.member_expenses or (bundle.expense,))
                if str(expense["invoice_buyer"] or "").strip()
            ),
            "",
        )
        note = next(
            (
                str(expense["substitute_reason"] or expense["note"] or "").strip()
                for expense in (bundle.member_expenses or (bundle.expense,))
                if str(expense["substitute_reason"] or expense["note"] or "").strip()
            ),
            "",
        )
        sheet.append(
            [
                bundle.group_id,
                company_alias(bundle.expense["company_entity"]),
                buyer,
                bundle.expense["employee_name"],
                bundle.expense_name,  # category: 差旅交通 / AI 项目 / …
                _ticket_status(bundle),
                len(bundle.invoices),
                invoice_total,
                actual_amount,
                round(invoice_total - actual_amount, 2),
                _attachment_count(bundle),
                _compose_remark(note, _bundle_pending_approval(bundle)),
            ]
        )


def _fill_invoices(sheet, bundles: list[ExportBundle]) -> None:
    sheet.append(
        [
            "公司主体",
            "购买方名称",
            "人员",
            "编号",
            "报销事项",
            "发票/单据类型",
            "发票号码",
            "开票日期",
            "销售方",
            "项目名称",
            "报销价税合计",
            "是否查验",
            "发票文件路径",
            "备注",
        ]
    )
    for bundle in bundles:
        for invoice in bundle.invoices:
            allocation = invoice.allocation
            item = invoice.invoice_item
            sheet.append(
                [
                    company_alias(bundle.expense["company_entity"]),
                    allocation["invoice_buyer"],
                    bundle.expense["employee_name"],
                    bundle.group_id,
                    invoice.project_name or bundle.expense_name,
                    allocation["invoice_type"] or _invoice_text(item, "sub_type_description", "type_description"),
                    allocation["invoice_number"],
                    allocation["invoice_date"],
                    _invoice_text(item, "seller", "seller_name", "seller_title"),
                    _invoice_text(
                        item,
                        "item_name",
                        "goods_name",
                        "service_name",
                        "project_name",
                        "sub_type_description",
                        "type_description",
                    ),
                    allocation["invoice_amount"],
                    "是" if str(allocation["invoice_number"] or "").strip() else "否",
                    invoice.document.package_path,
                    _compose_remark(
                        allocation["note"] or bundle.expense["substitute_reason"] or bundle.expense["note"] or "",
                        _is_pending_approval(invoice),
                    ),
                ]
            )


def _invoice_package_leaf(document: ExportDocument) -> str:
    """Innermost invoice image/file name from the package path (e.g. AI项目发票.pdf)."""
    if document.package_path:
        return Path(document.package_path).name
    return _source_leaf_filename(document)


def _person_company_entity(expense: sqlite3.Row) -> str:
    """Full company entity from the user's personal profile."""
    keys = set(expense.keys())
    if "user_company_entity" in keys and str(expense["user_company_entity"] or "").strip():
        return str(expense["user_company_entity"]).strip()
    return str(expense["company_entity"] or "").strip()


def _fill_attachments(sheet, bundles: list[ExportBundle]) -> None:
    sheet.append(["组ID", "公司主体", "人员", "附件类型", "文件名", "文件金额参考", "源文件路径", "说明"])
    for bundle in bundles:
        company = _person_company_entity(bundle.expense)
        evidence_count = len(bundle.transaction_attachments)
        if bundle.invoices:
            for invoice in bundle.invoices:
                invoice_path = (
                    invoice.document.package_path
                    if invoice.document.exists and invoice.document.package_path
                    else str(invoice.document.source_path)
                )
                invoice_name = _invoice_package_leaf(invoice.document)
                if not invoice.document.exists:
                    sheet.append(
                        [
                            bundle.group_id,
                            company,
                            bundle.expense["employee_name"],
                            "发票文件缺失",
                            invoice_name,
                            invoice.allocation["invoice_amount"],
                            invoice_path,
                            "数据库有发票记录，但源文件不存在",
                        ]
                    )
                    continue
                note = "发票图片，供核对使用"
                if evidence_count:
                    note = f"发票图片，供核对使用（同组另有 {evidence_count} 份佐证材料）"
                sheet.append(
                    [
                        bundle.group_id,
                        company,
                        bundle.expense["employee_name"],
                        "发票图片",
                        invoice_name,
                        invoice.allocation["invoice_amount"],
                        invoice_path,
                        _compose_remark(note, _is_pending_approval(invoice)),
                    ]
                )
            continue

        # No invoice: keep evidence rows for manual check, mark as 待核对.
        for document in bundle.transaction_attachments:
            sheet.append(
                [
                    bundle.group_id,
                    company,
                    bundle.expense["employee_name"],
                    document.attachment_kind,
                    _source_leaf_filename(document),
                    document.amount_hint if document.amount_hint is not None else _actual_total(bundle),
                    document.package_path if document.exists else str(document.source_path),
                    "文件缺失" if not document.exists else "暂无匹配发票，佐证材料待核对",
                ]
            )
        sheet.append(
            [
                bundle.group_id,
                company,
                bundle.expense["employee_name"],
                "待核对",
                "",
                _actual_total(bundle),
                "",
                "该报销项没有发票匹配明细",
            ]
        )


def _ticket_status(bundle: ExportBundle) -> str:
    members = bundle.member_expenses or (bundle.expense,)
    if any(member["is_substitute"] for member in members):
        return "替票"
    if not bundle.invoices:
        return "待核对"
    if bundle.transaction_attachments:
        return "正常发票/行程单"
    return "正常发票"


def _style_sheet(sheet) -> None:
    if sheet.title == "总览":
        _style_overview_sheet(sheet)
        return

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2E0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    max_row = sheet.max_row
    max_column = sheet.max_column
    if max_row == 0 or max_column == 0:
        return

    header_row = 1
    if max_row >= header_row:
        for cell in sheet[header_row]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in sheet.iter_rows():
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, (int, float)):
                cell.number_format = '#,##0.00'

    red_font = Font(color="FF0000")
    # Highlight 本次报销金额 column in summary sheet.
    if sheet.title == "报销项汇总" and max_column >= 9:
        for row in range(2, max_row + 1):
            sheet.cell(row=row, column=9).font = red_font

    # Mark「等待通过」remarks in red (备注 / 说明 columns).
    remark_columns = {
        "报销项汇总": 12,
        "发票明细": 14,
        "附件与待核对": 8,
    }
    remark_col = remark_columns.get(sheet.title)
    if remark_col:
        for row in range(2, max_row + 1):
            cell = sheet.cell(row=row, column=remark_col)
            if PENDING_APPROVAL_NOTE in str(cell.value or ""):
                cell.font = red_font

    for column_index in range(1, max_column + 1):
        width = 12
        column_letter = get_column_letter(column_index)
        for cell in sheet[column_letter]:
            value = "" if cell.value is None else str(cell.value)
            width = max(width, min(len(value) + 2, 36))
        sheet.column_dimensions[column_letter].width = width

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def _style_overview_sheet(sheet) -> None:
    navy = "1F4E78"
    white = "FFFFFF"
    pale_header = "F3F7FB"
    line_blue = "2BAAE2"
    red = "FF0000"
    border_side = Side(style="thin", color=line_blue)
    light_side = Side(style="thin", color="E5E7EB")
    currency_format = '"¥"#,##0.00;[Red]-"¥"#,##0.00'

    sheet.sheet_view.showGridLines = True
    sheet.freeze_panes = "A7"
    widths = {"A": 18, "B": 18, "C": 14, "D": 14, "E": 18, "F": 18, "G": 18, "I": 12}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, min_col=1, max_col=7):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = Border(left=light_side, right=light_side, top=light_side, bottom=light_side)
            if isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0"

    title = sheet["A1"]
    title.fill = PatternFill("solid", fgColor=navy)
    title.font = Font(color="FFFFFF", bold=True, size=18)
    title.alignment = Alignment(horizontal="center", vertical="center")
    for row in range(1, 3):
        sheet.row_dimensions[row].height = 28
        for column in range(1, 8):
            sheet.cell(row=row, column=column).fill = PatternFill("solid", fgColor=navy)

    for label_cell, value_cell in (("A4", "A5"), ("C4", "C5"), ("E4", "E5"), ("G4", "G5")):
        sheet[label_cell].fill = PatternFill("solid", fgColor=pale_header)
        sheet[label_cell].font = Font(color=navy, bold=True, size=12)
        sheet[label_cell].alignment = Alignment(horizontal="center")
        sheet[value_cell].font = Font(color=navy, bold=True, size=15)
        sheet[value_cell].alignment = Alignment(horizontal="center")
    sheet["A5"].number_format = currency_format
    sheet["C5"].number_format = currency_format

    for row_index in range(1, sheet.max_row + 1):
        first_value = sheet.cell(row=row_index, column=1).value
        if first_value == "公司主体":
            for column in range(1, 8):
                cell = sheet.cell(row=row_index, column=column)
                cell.fill = PatternFill("solid", fgColor=navy)
                cell.font = Font(color=red if column == 6 else "FFFFFF", bold=True, size=12)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
            _style_overview_table_body(sheet, row_index + 1, currency_format, white, line_blue, red)
        if first_value == "核对提示":
            for column in range(1, 8):
                cell = sheet.cell(row=row_index, column=column)
                cell.fill = PatternFill("solid", fgColor=navy)
                cell.font = Font(color="FFFFFF", bold=True, size=12)
                cell.alignment = Alignment(horizontal="center")


def _style_overview_table_body(sheet, start_row: int, currency_format: str, fill_color: str, line_color: str, red: str) -> None:
    side = Side(style="thin", color=line_color)
    row = start_row
    while row <= sheet.max_row and sheet.cell(row=row, column=1).value not in (None, "核对提示", "公司主体"):
        for column in range(1, 8):
            cell = sheet.cell(row=row, column=column)
            cell.fill = PatternFill("solid", fgColor=fill_color)
            cell.border = Border(left=side, right=side, top=side, bottom=side)
            if column in (5, 6, 7) and isinstance(cell.value, (int, float)):
                cell.number_format = currency_format
            if column == 6:
                cell.font = Font(color=red)
            cell.alignment = Alignment(horizontal="right" if column >= 2 else "left", vertical="center")
        row += 1
